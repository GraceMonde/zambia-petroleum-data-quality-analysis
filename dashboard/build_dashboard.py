from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.marker import Marker
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
CLEAN_PATH = ROOT / "data" / "cleaned_data.csv"
ISSUES_PATH = ROOT / "data" / "issues_flagged.csv"
ROUTES_PATH = ROOT / "data" / "route_metrics.csv"
OUTPUT_PATH = ROOT / "dashboard" / "ERB_transport_dashboard.xlsx"

NAVY = "1F3864"
LIGHT_BLUE = "D9E2F3"
PALE = "F2F2F2"
RED_FILL = "FCE4E4"
AMBER_FILL = "FFF2CC"
WHITE = "FFFFFF"
GREY_TEXT = "595959"

THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

CHART_WIDTH = 11.5
CHART_HEIGHT = 7.0
COLUMN_WIDTH = 13.5


def load_inputs():
    clean = pd.read_csv(CLEAN_PATH)
    clean["date"] = pd.to_datetime(clean["shipment_date"])
    issues = pd.read_csv(ISSUES_PATH)
    routes = pd.read_csv(ROUTES_PATH)
    return clean, issues, routes


def product_volume(clean):
    frame = (
        clean.groupby("product")["volume_litres"]
        .sum()
        .reset_index()
        .sort_values("volume_litres", ascending=False)
        .rename(columns={"product": "label", "volume_litres": "value"})
    )
    return frame


def province_volume(clean):
    frame = (
        clean.groupby("province")["volume_litres"]
        .sum()
        .reset_index()
        .sort_values("volume_litres", ascending=False)
        .rename(columns={"province": "label", "volume_litres": "value"})
    )
    return frame


def destination_cost(clean):
    frame = (
        clean.groupby(["destination", "province"])["transport_cost"]
        .mean()
        .reset_index()
        .sort_values("transport_cost", ascending=False)
        .head(10)
    )
    frame["label"] = frame["destination"] + " (" + frame["province"] + ")"
    return frame[["label", "transport_cost"]].rename(
        columns={"transport_cost": "value"}
    )


def weekly_cost(clean):
    frame = clean.copy()
    frame["week"] = frame["date"].dt.to_period("W").dt.start_time
    grouped = frame.groupby("week")["transport_cost"].mean().reset_index()
    grouped["label"] = grouped["week"].dt.strftime("%d %b")
    return grouped[["label", "transport_cost"]].rename(
        columns={"transport_cost": "value"}
    )


def monthly_summary(clean):
    frame = clean.copy()
    frame["month"] = frame["date"].dt.strftime("%Y-%m")
    return (
        frame.groupby("month")
        .agg(volume=("volume_litres", "sum"), avg_cost=("transport_cost", "mean"))
        .reset_index()
        .rename(columns={"month": "label"})
    )


def issue_counts(issues):
    return (
        issues.groupby(["issue", "severity"])
        .size()
        .reset_index(name="value")
        .sort_values("value", ascending=False)
    )


def build_kpis(clean, issues):
    return [
        ("Shipments analysed", len(clean), "#,##0"),
        ("Volume moved (litres)", int(clean["volume_litres"].sum()), "#,##0"),
        ("Avg transport cost (ZMW)", int(clean["transport_cost"].mean()), "#,##0"),
        ("Data quality issues", len(issues), "#,##0"),
    ]


def write_table(sheet, start_row, title, headers, rows, formats=None, start_column=1):
    formats = formats or {}
    sheet.cell(row=start_row, column=start_column, value=title).font = Font(
        bold=True, color=NAVY, size=11
    )
    header_row = start_row + 1
    for offset, header in enumerate(headers):
        cell = sheet.cell(
            row=header_row, column=start_column + offset, value=header
        )
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER

    for row_offset, row_values in enumerate(rows, start=1):
        for col_offset, value in enumerate(row_values):
            cell = sheet.cell(
                row=header_row + row_offset, column=start_column + col_offset, value=value
            )
            cell.border = BORDER
            if col_offset + 1 in formats:
                cell.number_format = formats[col_offset + 1]

    return {
        "header_row": header_row,
        "first_row": header_row + 1,
        "last_row": header_row + len(rows),
        "start_column": start_column,
    }


def build_data_sheet(workbook, clean, issues, routes):
    sheet = workbook.create_sheet("Data")
    sheet.sheet_view.showGridLines = False
    sheet["A1"] = (
        "Supporting tables - sources: data/cleaned_data.csv, data/issues_flagged.csv, "
        "data/route_metrics.csv"
    )
    sheet["A1"].font = Font(italic=True, color=GREY_TEXT)

    tables = [
        (
            "product",
            "Volume by product (clean records)",
            ["Product", "Volume (L)"],
            product_volume(clean)[["label", "value"]].values.tolist(),
            {2: "#,##0"},
        ),
        (
            "province",
            "Volume by destination province (clean records)",
            ["Province", "Volume (L)"],
            province_volume(clean)[["label", "value"]].values.tolist(),
            {2: "#,##0"},
        ),
        (
            "destination",
            "Average transport cost by destination - top 10 (clean records)",
            ["Destination", "Avg cost (ZMW)"],
            destination_cost(clean).values.tolist(),
            {2: "#,##0"},
        ),
        (
            "weekly",
            "Average transport cost by week (clean records)",
            ["Week commencing", "Avg cost (ZMW)"],
            weekly_cost(clean).values.tolist(),
            {2: "#,##0"},
        ),
        (
            "issues",
            "Data-quality issues by category",
            ["Issue", "Severity", "Records"],
            issue_counts(issues)[["issue", "severity", "value"]].values.tolist(),
            {3: "#,##0"},
        ),
        (
            "monthly",
            "Monthly summary (clean records)",
            ["Month", "Volume (L)", "Avg cost (ZMW)"],
            monthly_summary(clean)[["label", "volume", "avg_cost"]].values.tolist(),
            {2: "#,##0", 3: "#,##0"},
        ),
        (
            "flagged",
            "Routes flagged for further investigation",
            ["Route", "Cost/litre-km", "Vs network", "Anomalies", "Reason"],
            [
                [
                    row_values["route"],
                    row_values["cost_per_litre_km"],
                    row_values["deviation_pct"] / 100,
                    int(row_values["high_cost_anomalies"]),
                    row_values["flag_reason"],
                ]
                for _, row_values in routes[routes["flagged"]].iterrows()
            ],
            {2: "0.000000", 3: "+0.0%;-0.0%", 4: "0"},
        ),
    ]

    coords = {}
    row = 3
    for key, title, headers, rows, formats in tables:
        coords[key] = write_table(sheet, row, title, headers, rows, formats=formats)
        row = coords[key]["last_row"] + 3

    coords["register"] = write_table(
        sheet,
        3,
        "Data-quality exception register",
        ["Record", "Shipment", "Route", "Issue", "Severity", "Detail"],
        issues[
            ["record_id", "shipment_id", "route", "issue", "severity", "detail"]
        ].values.tolist(),
        formats={1: "0"},
        start_column=8,
    )

    widths = {1: 28, 2: 18, 3: 16, 4: 14, 5: 60, 6: 16}
    widths.update({8: 10, 9: 12, 10: 30, 11: 26, 12: 12, 13: 70})
    for column, width in widths.items():
        sheet.column_dimensions[get_column_letter(column)].width = width

    return sheet, coords


def style_chart(chart, title, category_title, value_title):
    chart.title = title
    chart.style = 10
    chart.width = CHART_WIDTH
    chart.height = CHART_HEIGHT
    chart.legend = None
    chart.x_axis.title = category_title
    chart.y_axis.title = value_title


def make_bar(sheet, coords, title, category_title, value_title, horizontal=False, labels=False, color="2E75B6", value_column=2):
    chart = BarChart()
    chart.type = "bar" if horizontal else "col"
    chart.gapWidth = 60
    chart.add_data(
        Reference(
            sheet,
            min_col=value_column,
            max_col=value_column,
            min_row=coords["header_row"],
            max_row=coords["last_row"],
        ),
        titles_from_data=True,
    )
    chart.set_categories(
        Reference(
            sheet, min_col=1, min_row=coords["first_row"], max_row=coords["last_row"]
        )
    )
    style_chart(chart, title, category_title, value_title)
    chart.series[0].graphicalProperties.solidFill = color
    chart.series[0].graphicalProperties.line.solidFill = color
    if labels:
        chart.dataLabels = DataLabelList()
        chart.dataLabels.showVal = True
    return chart


def make_line(sheet, coords, title, category_title, value_title, color="1F3864"):
    chart = LineChart()
    chart.add_data(
        Reference(
            sheet,
            min_col=2,
            min_row=coords["header_row"],
            max_row=coords["last_row"],
        ),
        titles_from_data=True,
    )
    chart.set_categories(
        Reference(
            sheet, min_col=1, min_row=coords["first_row"], max_row=coords["last_row"]
        )
    )
    style_chart(chart, title, category_title, value_title)
    series = chart.series[0]
    series.marker = Marker(symbol="circle", size=6)
    series.marker.graphicalProperties.solidFill = color
    series.marker.graphicalProperties.line.solidFill = color
    series.graphicalProperties.line.solidFill = color
    series.graphicalProperties.line.width = 22000
    series.smooth = False
    return chart


def style_kpi(dashboard, column, label, value, number_format):
    label_fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    value_fill = PatternFill("solid", fgColor=PALE)

    label_cell = dashboard.cell(row=4, column=column, value=label)
    label_cell.font = Font(bold=True, color=NAVY, size=10)
    label_cell.fill = label_fill
    label_cell.alignment = Alignment(horizontal="center", vertical="center")
    label_cell.border = BORDER

    value_cell = dashboard.cell(row=5, column=column, value=value)
    value_cell.font = Font(bold=True, size=20, color=NAVY)
    value_cell.fill = value_fill
    value_cell.alignment = Alignment(horizontal="center", vertical="center")
    value_cell.number_format = number_format
    value_cell.border = BORDER

    for row, fill in ((4, label_fill), (5, value_fill)):
        dashboard.merge_cells(
            start_row=row, start_column=column, end_row=row, end_column=column + 1
        )
        trailing = dashboard.cell(row=row, column=column + 1)
        trailing.border = BORDER
        trailing.fill = fill


def section_header(sheet, row, text, start_column=1, end_column=11):
    if start_column != 1 or end_column != start_column:
        sheet.merge_cells(
            start_row=row,
            start_column=start_column,
            end_row=row,
            end_column=end_column,
        )
    cell = sheet.cell(row=row, column=start_column, value=text)
    cell.font = Font(bold=True, size=12, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    sheet.row_dimensions[row].height = 20


def write_spanned_row(sheet, row, specs):
    for start_column, end_column, value, options in specs:
        cell = sheet.cell(row=row, column=start_column, value=value)
        cell.border = BORDER
        fill = options.get("fill")
        if fill:
            cell.fill = PatternFill("solid", fgColor=fill)
        if options.get("bold"):
            cell.font = Font(bold=True, color=WHITE if fill else "000000")
        if options.get("center"):
            cell.alignment = Alignment(horizontal="center", vertical="center")
        if options.get("number_format"):
            cell.number_format = options["number_format"]
        if end_column > start_column:
            sheet.merge_cells(
                start_row=row,
                start_column=start_column,
                end_row=row,
                end_column=end_column,
            )
            for column in range(start_column + 1, end_column + 1):
                extra = sheet.cell(row=row, column=column)
                extra.border = BORDER
                if fill:
                    extra.fill = PatternFill("solid", fgColor=fill)


def write_dashboard(workbook, clean, issues, routes, coords, data_sheet):
    dashboard = workbook.active
    dashboard.title = "Dashboard"
    dashboard.sheet_view.showGridLines = False

    for column in range(1, 14):
        dashboard.column_dimensions[get_column_letter(column)].width = COLUMN_WIDTH

    dashboard.merge_cells("A1:L1")
    title = dashboard["A1"]
    title.value = "Fuel Transport Data Quality & Analysis - Zambia"
    title.font = Font(bold=True, size=18, color=WHITE)
    title.fill = PatternFill("solid", fgColor=NAVY)
    title.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    dashboard.row_dimensions[1].height = 32

    dashboard.merge_cells("A2:L2")
    subtitle = dashboard["A2"]
    subtitle.value = (
        "SYNTHETIC PORTFOLIO DATASET - fictional petroleum product movements. "
        "Not an ERB dataset and not a representation of the UPP model."
    )
    subtitle.font = Font(italic=True, size=10, color="C00000")
    subtitle.alignment = Alignment(horizontal="left", vertical="center", indent=1)

    for position, (label, value, number_format) in enumerate(build_kpis(clean, issues)):
        style_kpi(dashboard, 1 + position * 3, label, value, number_format)

    total_received = len(clean) + len(issues)
    context = dashboard.cell(
        row=6,
        column=1,
        value=(
            f"{total_received} records received  |  {len(clean)} analysed "
            f"({len(clean) / total_received * 100:.1f}% clean)  |  {len(issues)} "
            "records excluded from analysis and held for review  |  "
            "Source: data/cleaned_data.csv and data/issues_flagged.csv"
        ),
    )
    context.font = Font(italic=True, size=9, color=GREY_TEXT)
    context.alignment = Alignment(vertical="center")
    dashboard.merge_cells("A6:L6")

    charts = [
        ("A8", make_bar(data_sheet, coords["product"], "Petroleum volume by product", "Product", "Volume (L)", labels=True)),
        ("F8", make_bar(data_sheet, coords["province"], "Volume by destination region", "Province", "Volume (L)", horizontal=True)),
        ("A25", make_bar(data_sheet, coords["destination"], "Average transport cost by destination (top 10)", "Destination", "Average cost (ZMW)", horizontal=True, color="ED7D31")),
        ("F25", make_line(data_sheet, coords["weekly"], "Transport cost trend over time (weekly average)", "Week commencing", "Average cost (ZMW)")),
        ("A42", make_bar(data_sheet, coords["issues"], "Data-quality issues by category", "Issue", "Records", horizontal=True, labels=True, color="C00000", value_column=3)),
    ]
    for anchor, chart in charts:
        dashboard.add_chart(chart, anchor)

    flagged = routes[routes["flagged"]].sort_values("cost_per_litre_km", ascending=False)
    section_header(dashboard, 41, "Routes flagged for further investigation", 6, 11)
    write_spanned_row(
        dashboard,
        42,
        [
            (6, 7, "Route", {"fill": NAVY, "bold": True, "center": True}),
            (8, 8, "Cost/litre-km", {"fill": NAVY, "bold": True, "center": True}),
            (9, 9, "Vs network", {"fill": NAVY, "bold": True, "center": True}),
            (10, 10, "Anomalies", {"fill": NAVY, "bold": True, "center": True}),
        ],
    )
    for offset, (_, row_values) in enumerate(flagged.iterrows(), start=43):
        write_spanned_row(
            dashboard,
            offset,
            [
                (6, 7, row_values["route"], {"fill": AMBER_FILL}),
                (8, 8, row_values["cost_per_litre_km"], {"fill": AMBER_FILL, "number_format": "0.000000", "center": True}),
                (9, 9, row_values["deviation_pct"] / 100, {"fill": AMBER_FILL, "number_format": "+0.0%;-0.0%", "center": True}),
                (10, 10, int(row_values["high_cost_anomalies"]), {"fill": AMBER_FILL, "center": True}),
            ],
        )

    register_title_row = 59
    section_header(dashboard, register_title_row, "Data-quality exception register", 1, 13)
    write_spanned_row(
        dashboard,
        register_title_row + 1,
        [
            (1, 1, "Record", {"fill": NAVY, "bold": True, "center": True}),
            (2, 2, "Shipment", {"fill": NAVY, "bold": True, "center": True}),
            (3, 4, "Route", {"fill": NAVY, "bold": True, "center": True}),
            (5, 6, "Issue", {"fill": NAVY, "bold": True, "center": True}),
            (7, 7, "Severity", {"fill": NAVY, "bold": True, "center": True}),
        ],
    )
    for offset, (_, row_values) in enumerate(issues.iterrows(), start=register_title_row + 2):
        fill = RED_FILL if row_values["severity"] == "High" else AMBER_FILL
        write_spanned_row(
            dashboard,
            offset,
            [
                (1, 1, int(row_values["record_id"]), {"fill": fill, "center": True}),
                (2, 2, row_values["shipment_id"], {"fill": fill, "center": True}),
                (3, 4, row_values["route"], {"fill": fill}),
                (5, 6, row_values["issue"], {"fill": fill}),
                (7, 7, row_values["severity"], {"fill": fill, "center": True}),
            ],
        )

    note_row = register_title_row + len(issues) + 3
    section_header(dashboard, note_row, "Method and sources", 1, 13)
    notes = [
        "Collection, cleaning, validation and flagging: Python/pandas - src/data_cleaning.py. "
        "Validation rules and counts: reports/data_quality_report.md.",
        "SQL analysis: SQLite - sql/analysis_queries.sql, with results in sql/query_results.md. "
        "Route cost analysis: src/analysis.py and reports/analytical_summary.md.",
        "Severity: High = hard validation failure (missing, invalid or out-of-tolerance value); "
        "Medium = duplicate or statistical flag requiring review.",
        "Charts are driven by the Data sheet; refresh by re-running "
        "dashboard/build_dashboard.py after any data refresh.",
    ]
    for offset, note in enumerate(notes, start=1):
        row = note_row + offset
        dashboard.merge_cells(start_row=row, start_column=1, end_row=row, end_column=13)
        cell = dashboard.cell(row=row, column=1, value=note)
        cell.font = Font(size=9, color=GREY_TEXT)
        cell.alignment = Alignment(vertical="center", wrap_text=False)

    dashboard.freeze_panes = "A8"


def main():
    clean, issues, routes = load_inputs()
    workbook = Workbook()
    data_sheet, coords = build_data_sheet(workbook, clean, issues, routes)
    write_dashboard(workbook, clean, issues, routes, coords, data_sheet)
    workbook.save(OUTPUT_PATH)
    print(f"wrote {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
