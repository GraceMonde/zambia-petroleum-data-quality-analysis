import re
import sqlite3
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
failures = []


def check(label, condition, detail=""):
    if condition:
        print(f"  PASS  {label}")
    else:
        failures.append(label)
        print(f"  FAIL  {label} {detail}")


expected_files = [
    "README.md",
    "requirements.txt",
    "data/raw_data.csv",
    "data/processed_data.csv",
    "data/cleaned_data.csv",
    "data/issues_flagged.csv",
    "data/route_metrics.csv",
    "data/warehouse.db",
    "src/generate_data.py",
    "src/data_cleaning.py",
    "src/load_database.py",
    "src/run_queries.py",
    "src/analysis.py",
    "sql/analysis_queries.sql",
    "sql/query_results.md",
    "dashboard/build_dashboard.py",
    "dashboard/ERB_transport_dashboard.xlsx",
    "dashboard/powerbi_build_spec.md",
    "reports/data_quality_report.md",
    "reports/analytical_summary.md",
    "docs/guide.txt",
    "docs/project_report.md",
    "docs/screenshots/.gitkeep",
    "tests/verify_project.py",
]
print("files")
for name in expected_files:
    check(name, (ROOT / name).exists())

raw = pd.read_csv(ROOT / "data/raw_data.csv")
clean = pd.read_csv(ROOT / "data/cleaned_data.csv")
processed = pd.read_csv(ROOT / "data/processed_data.csv")
issues = pd.read_csv(ROOT / "data/issues_flagged.csv")
routes = pd.read_csv(ROOT / "data/route_metrics.csv")

print("\ndataset")
check("raw rows = 408", len(raw) == 408, len(raw))
check("processed rows = 408", len(processed) == 408, len(processed))
check("clean rows = 353", len(clean) == 353, len(clean))
check("issue rows = 55", len(issues) == 55, len(issues))
check("raw has surrogate key", raw["record_id"].is_unique)
check(
    "8 duplicate business keys in raw",
    int(raw["shipment_id"].duplicated().sum()) == 8,
    int(raw["shipment_id"].duplicated().sum()),
)
check("no duplicate record_id", not issues["record_id"].duplicated().any())

expected_counts = {
    "DISTANCE_MISMATCH": 17,
    "MISSING_DISTANCE": 11,
    "HIGH_COST": 9,
    "DUPLICATE": 8,
    "MISSING_DIMENSION": 6,
    "INVALID_VOLUME": 4,
    "DATE_INVALID": 0,
}
counts = issues["issue_code"].value_counts().to_dict()
print("\nissue counts")
for code, expected in expected_counts.items():
    actual = counts.get(code, 0)
    check(f"{code} = {expected}", actual == expected, actual)

print("\ncleaning invariants")
check("clean volumes all positive", bool((clean["volume_litres"] > 0).all()))
check("clean has no nulls", int(clean.isna().sum().sum()) == 0)
check(
    "clean claimed == reference",
    bool((clean["claimed_distance_km"] == clean["reference_distance_km"]).all()),
)
check(
    "products standardised",
    set(clean["product"]) == {"Petrol", "Diesel", "Kerosene", "Jet A-1"},
    sorted(set(clean["product"])),
)
check("clean record_ids unique", clean["record_id"].is_unique)
check(
    "flagged records excluded from clean",
    len(set(clean["record_id"]) & set(issues["record_id"])) == 0,
)
check(
    "clean + flagged = total",
    len(clean) + issues["record_id"].nunique() == 408,
)
check(
    "status column correct",
    int((processed["quality_status"] == "flagged").sum()) == 55,
)

print("\nanalysis")
check("26 routes", len(routes) == 26, len(routes))
check("3 routes flagged", int(routes["flagged"].sum()) == 3)
flagged_routes = sorted(routes[routes["flagged"]]["route"])
check(
    "flagged routes expected",
    flagged_routes
    == ["Livingstone -> Kazungula", "Lusaka -> Chipata", "Lusaka -> Kitwe"],
    flagged_routes,
)
summary = (ROOT / "reports/analytical_summary.md").read_text(encoding="utf-8")
check("summary states 0.004431 network rate", "0.004431" in summary)
check("summary lists 3 flagged routes", all(r in summary for r in flagged_routes))

print("\ndatabase")
connection = sqlite3.connect(ROOT / "data/warehouse.db")
for table, expected in [
    ("shipments", 408),
    ("shipment_issues", 55),
    ("companies", 7),
    ("depots", 7),
    ("destinations", 21),
    ("routes", 26),
]:
    actual = connection.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0]
    check(f"{table} = {expected}", actual == expected, actual)
check(
    "no foreign key violations",
    not connection.execute("PRAGMA foreign_key_check").fetchall(),
)
clean_volume = connection.execute(
    "SELECT CAST(SUM(volume_litres) AS INTEGER) FROM shipments WHERE quality_status='clean'"
).fetchone()[0]
check("clean volume 11,562,000", clean_volume == 11_562_000, clean_volume)
connection.close()

print("\nsql results")
results = (ROOT / "sql/query_results.md").read_text(encoding="utf-8")
check("9 queries executed", results.count("row(s) returned") == 9)
for token in ["408", "353", "86.5", "11,562,000", "17", "Distance discrepancy"]:
    check(f"results contain {token!r}", token in results)

print("\ndashboard")
workbook = load_workbook(ROOT / "dashboard/ERB_transport_dashboard.xlsx")
dashboard = workbook["Dashboard"]
check("sheets = Dashboard, Data", workbook.sheetnames == ["Dashboard", "Data"])
check("5 charts", len(dashboard._charts) == 5, len(dashboard._charts))
kpis = [dashboard.cell(row=5, column=c).value for c in (1, 4, 7, 10)]
check(
    "KPI values",
    kpis[0] == 353 and kpis[1] == 11_562_000 and kpis[3] == 55 and 39_000 < kpis[2] < 40_000,
    kpis,
)
register_rows = sum(
    1
    for row in range(61, 200)
    if isinstance(dashboard.cell(row=row, column=1).value, int)
)
check("exception register 55 rows", register_rows == 55, register_rows)
for chart in dashboard._charts:
    series = chart.series[0]
    reference = series.val.numRef.f
    check(f"chart references Data sheet: {reference}", reference.startswith("'Data'!"))
data_sheet = workbook["Data"]
check(
    "data sheet holds flagged routes",
    data_sheet["A75"].value in flagged_routes,
    data_sheet["A75"].value,
)

print("\nreadme claims")
readme = (ROOT / "README.md").read_text(encoding="utf-8")
for token in [
    "353",
    "408",
    "86.5%",
    "11,562,000",
    "0.004431",
    "17",
    "9 queries",
    "portfolio/synthetic dataset",
    "Lusaka -> Chipata",
    "requirements.txt",
]:
    check(f"README contains {token!r}", token in readme)

print("\n" + ("ALL CHECKS PASSED" if not failures else f"{len(failures)} FAILURES: {failures}"))
sys.exit(1 if failures else 0)
